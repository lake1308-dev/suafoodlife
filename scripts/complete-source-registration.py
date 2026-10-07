"""Append reviewed imported products and eligible RDA records, preserving nulls."""
import gzip, hashlib, json, math, sys
from pathlib import Path
import openpyxl
root = Path(__file__).resolve().parents[1]
fields = {'kcal': '에너지(kcal)', 'protein_g': '단백질(g)', 'fat_g': '지방(g)', 'carbs_g': '탄수화물(g)', 'sugars_g': '당류(g)', 'fiber_g': '식이섬유(g)', 'sodium_mg': '나트륨(mg)', 'cholesterol_mg': '콜레스테롤(mg)', 'sat_fat_g': '포화지방산(g)'}
def num(v):
    if v is None or str(v).strip() in ('', '-', 'Tr', 'tr', 'ND', 'N.D.'): return None
    try: value = float(v)
    except (ValueError, TypeError): return None
    assert math.isfinite(value) and value >= 0
    return value
def clean(v):
    return '' if v is None or str(v).strip() == '해당없음' else str(v).strip()
def write(path, data):
    encoded = gzip.compress(json.dumps(data, ensure_ascii=False, separators=(',', ':')).encode(), mtime=0)
    temp = root / (path + '.tmp'); temp.write_bytes(encoded); temp.replace(root / path)
    return hashlib.sha256(encoded).hexdigest()

manifest = json.loads((root/'data/catalog-manifest.json').read_text())
assert manifest['record_count'] == 274732
review = json.load(gzip.open(root/'data/unregistered-products-review.json.gz'))
wanted = {x['code']: x for x in review['records']}
assert len(wanted) == 63048
source = Path(sys.argv[1])
assert hashlib.sha256(source.read_bytes()).hexdigest() == review['source_sha256']
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
it = wb.active.iter_rows(values_only=True); idx = {x:i for i,x in enumerate(next(it))}
items = []; checked = set()
for rownum, row in enumerate(it, 2):
    get = lambda name: row[idx[name]]
    code = clean(get('식품코드'))
    if code not in wanted: continue
    assert code not in checked
    expected = wanted[code]
    assert rownum == expected['source_row'] and get('수입여부') == 'Y'
    assert clean(get('식품명')) == expected['name']
    basis = clean(get('영양성분함량기준량'))
    assert basis == expected['basis'] and basis in ('100g', '100ml')
    item = {'code':code, 'name':expected['name'], 'type':'가공식품', 'basis':basis,
            **{key:num(get(header)) for key,header in fields.items()},
            'source':'식품영양성분 데이터베이스(K-FCDB) · 식품의약품안전처',
            'reference_date':'2026-09-29', 'manufacturer':clean(get('제조사명')),
            'declared_weight':clean(get('식품중량')), 'import_flag':'Y'}
    assert item['kcal'] is not None
    items.append(item); checked.add(code)
    if len(items)%20000 == 0: print('Imported product rows:',len(items),flush=True)
wb.close()
assert checked == set(wanted)
columns = ['code','name','type','basis',*fields,'source','reference_date','manufacturer','declared_weight','import_flag']
for i,start in enumerate(range(0,len(items),5000),1):
    subset = sorted(items,key=lambda x:x['code'])[start:start+5000]
    path = f'data/catalog-imported-20260929-{i}.json.gz'
    sha = write(path,{'record_count':len(subset),'columns':columns,'rows':[[x.get(k) for k in columns] for x in subset]})
    manifest['chunks'].append({'path':path,'record_count':len(subset),'sha256':sha})
manifest['record_count'] += len(items)
manifest['version'] = '2026-10-07-registration-complete'
(root/'data/catalog-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))

rda_source = Path(sys.argv[2])
rda_review = json.loads((root/'data/unregistered-rda-review.json').read_text())
assert hashlib.sha256(rda_source.read_bytes()).hexdigest() == rda_review['source_sha256']
eligible = {x['code']:x for x in rda_review['records'] if x['decision']=='reprocess'}
wb = openpyxl.load_workbook(rda_source,read_only=True,data_only=True)
lookup = {r[0]:r[1:5] for r in list(wb['부록2)식품코드,국문명,영문명,학명 정보 '].values)[1:]}
positions = {'kcal':5,'protein_g':7,'fat_g':8,'carbs_g':10,'sugars_g':11,'fiber_g':18,'sodium_mg':26,'cholesterol_mg':87,'sat_fat_g':90}
rda_items=[]
for rownum,row in enumerate(wb['국가표준식품성분 Database 10.4'].values,1):
    if rownum<4 or not isinstance(row[0],(int,float)): continue
    code,name,en,latin = lookup[int(row[0])]
    if code not in eligible: continue
    assert rownum==eligible[code]['source_row'] and name==row[3]
    nutrients={k:num(row[p]) for k,p in positions.items()}
    assert all(nutrients[k] is not None for k in ('kcal','protein_g','fat_g','carbs_g'))
    rda_items.append({'id':'rda:'+code,'names':{'ko':name,'en':en},'aliases':{'ko':[],'en':[]},'category':'official_food','verification_status':'verified','nutrition_per_100g':nutrients,'allergen_info':{'contains':[],'status':'unverified','basis':'composition_not_verified'},'sources':[{'source':'농촌진흥청 국가표준식품성분 DB 10.4 (2026)','source_record_id':code,'source_food_name':name,'source_url':'https://www.nics.go.kr/food/kfi/fct/fctIntro/list','basis':'100g','data_type':'공식 식품 자료 · 가식부 100g','reference_date':'2026-10-07','analysis_source':row[4],'workbook_sheet':'국가표준식품성분 Database 10.4','workbook_row':rownum}]})
wb.close()
assert len(rda_items)==811 and len({x['id'] for x in rda_items})==811
write('data/rda-foods-remaining-811.json.gz',{'record_count':811,'original_sha256':rda_review['source_sha256'],'curated_id_by_record':{},'ingredients':rda_items})
display = ['kcal','protein_g','fat_g','carbs_g','sugars_g','sodium_mg','cholesterol_mg','sat_fat_g']
report={'imported_products_added':63048,'rda_foods_added':811,'rda_source_rows_held':55,'catalog_count':manifest['record_count'],'rda_count':3311,'combined_count':manifest['record_count']+3311,'original_sources':{'products':review['source_sha256'],'rda':rda_review['source_sha256']},'new_products_complete_display':sum(all(x[k] is not None for k in display) for x in items),'new_rda_complete_display':sum(all(x['nutrition_per_100g'][k] is not None for k in display) for x in rda_items),'allergy_values_inferred':0,'policy':'Append exact-code official source records, mark imported products, preserve original units and missing nutrition, hold 55 RDA records with incomplete core nutrition.'}
(root/'data/registration-completion-review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
