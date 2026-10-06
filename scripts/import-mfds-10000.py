"""Import 10,000 additional domestic products from the supplied K-FCDB workbook."""
import sys, json, gzip, hashlib, math
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1])
existing = set()
for p in (root / 'data').glob('nutrition-bulk-*.json.gz'):
    d = json.load(gzip.open(p)); ci = d['columns'].index('code')
    if d.get('version') == '2026-09-29' and d.get('source_sha256'): continue
    existing.update(r[ci] for r in d['rows'])
assert len(existing) == 50000
columns = ['code','name','type','basis','kcal','protein_g','fat_g','carbs_g','sugars_g','fiber_g','sodium_mg','cholesterol_mg','sat_fat_g','source','reference_date']
fields = {'kcal':'에너지(kcal)','protein_g':'단백질(g)','fat_g':'지방(g)','carbs_g':'탄수화물(g)','sugars_g':'당류(g)','fiber_g':'식이섬유(g)','sodium_mg':'나트륨(mg)','cholesterol_mg':'콜레스테롤(mg)','sat_fat_g':'포화지방산(g)'}
def number(v):
    if v is None or str(v).strip() in ('','-','Tr','ND'): return None
    try: n = float(v)
    except (ValueError,TypeError): return None
    return n if math.isfinite(n) and n >= 0 else None
def clean(v):
    s = str(v or '').strip()
    return '' if s == '해당없음' else s
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
ws = wb.active
it = ws.iter_rows(values_only=True)
headers = next(it); idx = {h:i for i,h in enumerate(headers)}
rows=[]; details={}; original=[]; seen=set(existing); scanned=0
for rownum, r in enumerate(it,2):
    scanned += 1
    get = lambda k: r[idx[k]]
    code = clean(get('식품코드'))
    if not code or code in seen or get('수입여부') != 'N' or get('데이터구분명') != '가공식품': continue
    basis = clean(get('영양성분함량기준량'))
    if basis not in ('100g','100ml'): continue
    values = {k:number(get(v)) for k,v in fields.items()}
    if any(values[k] is None for k in ('kcal','protein_g','fat_g','carbs_g')): continue
    name = clean(get('식품명')).lstrip('\ufeff')
    if not name: continue
    out={'code':code,'name':name,'type':'가공식품','basis':basis,**values,'source':'식품영양성분 데이터베이스(K-FCDB) · 식품의약품안전처','reference_date':'2026-09-29'}
    rows.append([out[c] for c in columns]); seen.add(code)
    details[code]={'manufacturer':clean(get('제조사명')),'declared_weight':clean(get('식품중량'))}
    original.append({'code':code,'worksheet_row':rownum})
    if len(rows)%1000 == 0: print('Selected',len(rows),'scanned',scanned,flush=True)
    if len(rows)==10000: break
assert len(rows)==10000, len(rows)
assert len(seen)==60000 and not (set(details)&existing)
sha=hashlib.sha256(source.read_bytes()).hexdigest()
for i in range(10):
    subset=rows[i*1000:(i+1)*1000]
    d={'version':'2026-09-29','source':'K-FCDB / 식품의약품안전처','source_sha256':sha,'record_count':1000,'columns':columns,'rows':subset}
    with gzip.GzipFile(filename=str(root/'data'/f'nutrition-bulk-{50001+i*1000}-{51000+i*1000}.json.gz'),mode='wb',mtime=0) as f: f.write(json.dumps(d,ensure_ascii=False,separators=(',',':')).encode())
for i in range(2):
    codes=[r[0] for r in rows[i*5000:(i+1)*5000]]
    with gzip.GzipFile(filename=str(root/'data'/f'food-details-{11+i}.json.gz'),mode='wb',mtime=0) as f: f.write(json.dumps({'version':'2026-09-29','source':'K-FCDB','records':{c:details[c] for c in codes}},ensure_ascii=False,separators=(',',':')).encode())
review={'source_filename':source.name,'source_sha256':sha,'source_rows':ws.max_row-1,'scanned_rows':scanned,'added_count':10000,'previous_count':50000,'result_count':60000,'duplicate_codes':0,'selection':'First 10,000 source-order records with import flag N, processed-food type, 100g/100ml basis, and all four core nutrients; exclude existing codes.','allergy_changes':0,'missing_values_preserved':sum(v is None for r in rows for v in r[4:13]),'records':original}
records=review.pop('records')
review['record_manifest']='mfds-import-records-10000.json.gz'
with gzip.GzipFile(filename=str(root/'data'/review['record_manifest']),mode='wb',mtime=0) as f: f.write(json.dumps(records,separators=(',',':')).encode())
(root/'data'/'mfds-import-review-10000.json').write_text(json.dumps(review,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in review.items() if k!='records'},ensure_ascii=False),flush=True)
